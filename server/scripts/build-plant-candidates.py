# /// script
# requires-python = ">=3.13"
# dependencies = ["pdfplumber==0.11.7", "openpyxl==3.1.5"]
# ///
"""Extract a provenance-bearing REVIEW pool, never a training-approved species list."""

import argparse
import csv
import hashlib
import json
import re
import sqlite3
from collections import Counter, defaultdict
from contextlib import closing
from datetime import UTC, datetime
from pathlib import Path

ROOT = Path(__file__).resolve().parents[2]
CACHE = ROOT / "server/.cache/research/common-plant-sources-20260922"
REGISTRY = ROOT / "docs/research/china-everyday-1000-sources.json"
CATALOG = ROOT / "server/.cache/catalog/treeoflife-10m/catalog.sqlite3"
SEEDS = ROOT / "docs/research/china-everyday-plants-candidates.json"
PDF_CONFIG = {
  "F18": {"rows": 114, "columns": 13, "start": 0, "flags": 5,
          "province": "黑龙江", "region": "东北", "work": "heilongjiang-native-2021",
          "uses": ["荒山", "平原", "城市", "乡村", "通道", "采矿迹地", "水系"]},
  "F19": {"rows": 163, "columns": 11, "start": 0, "flags": 4,
          "province": "湖北", "region": "华中", "work": "hubei-native-2020",
          "uses": ["荒山", "平原", "城市", "乡村", "通道", "水系"]},
  "F20": {"rows": 140, "columns": 12, "start": 2, "flags": 5,
          "province": "北京", "region": "华北", "work": "beijing-native-2021",
          "uses": ["荒山", "平原", "城市", "乡村", "通道", "水系"]},
}
FREQUENCIES = ("极常见", "常见", "多见", "极少见", "少见", "罕见", "偶见", "近年未见",
               "近年未采到", "标本未见")
STRONG_FREQUENCIES = {"极常见", "常见", "多见"}


def digest(path):
  with path.open("rb") as stream:
    return hashlib.file_digest(stream, "sha256").hexdigest()


def compact(value):
  return re.sub(r"\s+", "", value or "")


def text(value):
  return re.sub(r"\s+", " ", str(value or "")).strip()


def write_json(path, value):
  path.write_text(json.dumps(value, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")


def frequency(note):
  """Taxonomic '好种' is deliberately NOT evidence of commonness."""
  first_clause = re.split("[，,；;。]", text(note))[0]
  return first_clause if first_clause in FREQUENCIES else "未明确"


def checked_file(cache, source):
  filename = source.get("localFile", f"{source['id']}.pdf")
  path = cache / filename
  if digest(path) != source["directFetchCheck"]["sha256"]:
    raise ValueError(f"{source['id']}: source checksum mismatch.")
  return path


def extract_pdf(cache, source, config):
  import pdfplumber

  path = checked_file(cache, source)
  records = []
  with pdfplumber.open(path) as pdf:
    fixed_x = None
    if source["id"] == "F18":
      # A page-6 table has underlined text that looks like extra cell borders.
      # Fix only the verified column grid; real row boundaries come from column 1.
      table = pdf.pages[0].find_tables()[0]
      fixed_x = sorted({round(col.bbox[0], 2) for col in table.columns}
                       | {round(table.bbox[2], 2)})
      if len(fixed_x) != config["columns"] + 1:
        raise ValueError("Unexpected Heilongjiang grid.")
    for page in pdf.pages[config["start"]:]:
      settings = {}
      if fixed_x is not None:
        ys = sorted({round(edge["top"], 2) for edge in page.edges
                     if edge["orientation"] == "h"
                     and edge["x0"] <= fixed_x[0] + 1 and edge["x1"] >= fixed_x[1] - 1})
        settings = {
          "vertical_strategy": "explicit", "explicit_vertical_lines": fixed_x,
          "horizontal_strategy": "explicit", "explicit_horizontal_lines": ys,
        }
      for table in page.find_tables(settings):
        for cells in table.extract(x_tolerance=1):
          if len(cells) != config["columns"]:
            raise ValueError(f"{source['id']} page {page.page_number}: changed column grid.")
          first = text(cells[0])
          if "序" in first or any("拉丁学名" in (c or "") for c in cells):
            continue
          if not first.isdigit():
            # Only observed continuation in these sources: author split over a page.
            nonempty = {i for i, cell in enumerate(cells) if cell}
            if nonempty == {2} and records:
              records[-1]["scientificNameRaw"] += "\n" + cells[2]
              records[-1]["pages"].append(page.page_number)
            elif any("绿化" in compact(c) for c in cells if c):
              continue  # second header row
            elif nonempty:
              raise ValueError(f"{source['id']} page {page.page_number}: unknown continuation.")
            continue
          index = int(first)
          uses = []
          for label, value in zip(
            config["uses"], cells[config["flags"]:config["flags"] + len(config["uses"])],
            strict=True,
          ):
            if compact(value) not in {"", "○", "◯", "√"}:
              raise ValueError(f"{source['id']} row {index}: unknown suitability marker.")
            if compact(value):
              uses.append(label)
          records.append({
            "recordId": f"{source['id']}:{index}", "sourceId": source["id"],
            "sourceWork": config["work"], "sourceRow": index, "pages": [page.page_number],
            "sourceProvince": config["province"], "sourceRegion": config["region"],
            "chineseNameRaw": compact(cells[1]), "scientificNameRaw": cells[2] or "",
            "sourceRank": None, "sourceType": "native-recommendation",
            "suitableUses": uses, "suitableRegionRaw": compact(cells[-1]),
            "frequency": "未明确", "noteRaw": "",
            "evidence": "C-recommended-suitability-not-abundance",
          })
  if [r["sourceRow"] for r in records] != list(range(1, config["rows"] + 1)):
    raise ValueError(f"{source['id']}: missing, duplicated or out-of-order row IDs.")
  return records


def extract_workbook(cache, source):
  import openpyxl

  path = checked_file(cache, source)
  records = []
  with closing(openpyxl.load_workbook(path, read_only=True, data_only=True)) as workbook:
    for sheet_name, header, width, expected in [
      ("上海野生和逸生植物", 2, 28, 1238), ("上海栽培植物", 1, 21, 2888),
    ]:
      sheet = workbook[sheet_name]
      # The cultivated sheet falsely declares 16,379 columns because of formatting.
      # Never iterate using max_column; only the verified data columns are read.
      headers = next(sheet.iter_rows(
        min_row=header, max_row=header, max_col=width, values_only=True,
      ))
      for column, prefix in [(0, "序号"), (7, "中文名"), (8, "学名"),
                             (11, "等级"), (width - 1, "说明")]:
        if not str(headers[column]).startswith(prefix):
          raise ValueError(f"{sheet_name}: unexpected header.")
      wild = sheet_name == "上海野生和逸生植物"
      ids = []
      for row_number, row in enumerate(sheet.iter_rows(
        min_row=header + 1, max_col=width, values_only=True,
      ), header + 1):
        if row[0] is None:
          continue
        if not isinstance(row[0], int) or not row[7] or not row[8]:
          raise ValueError(f"{sheet_name}:{row_number}: malformed source row.")
        ids.append(row[0])
        label = frequency(row[-1])
        records.append({
          "recordId": f"{source['id']}:{'wild' if wild else 'cultivated'}:{row[0]}",
          "sourceId": source["id"], "sourceWork": source["sourceWork"],
          "sheet": sheet_name, "excelRow": row_number, "sourceRow": row[0],
          "sourceProvince": "上海", "sourceRegion": "华东",
          "chineseNameRaw": text(row[7]), "scientificNameRaw": text(row[8]),
          "sourceAuthor": text(row[9]), "sourceRank": text(row[11]),
          "sourceType": "survey-wild" if wild else "survey-cultivated",
          "growthStatusRaw": text(row[15 if wild else 13]),
          "taxonomyAssessmentRaw": text(row[19 if wild else 14]),
          "frequency": label, "noteRaw": text(row[-1]), "suitableUses": [],
          "evidence": (
            "A-regional-survey-common" if label in {"常见", "极常见"}
            else "A-regional-survey-frequent" if label == "多见"
            else "D-presence-or-rare-not-commonness"
          ),
        })
      if ids != list(range(1, expected + 1)):
        raise ValueError(f"{sheet_name}: changed count or missing row IDs.")
  return records


def source_html_body(cache, source):
  from html.parser import HTMLParser

  class Body(HTMLParser):
    def __init__(self):
      super().__init__()
      self.parts = []

    def handle_data(self, value):
      if value.strip():
        self.parts.append(value.strip())

  path = cache / f"{source['id']}.html"
  if digest(path) != source["directFetchCheck"]["sha256"]:
    raise ValueError("HTML source checksum mismatch.")
  body = Body()
  body.feed(path.read_text(encoding="utf-8"))
  return "\n".join(body.parts)


def extract_horticulture(cache, source):
  content = source_html_body(cache, source)
  match = re.search(r"——增加了(金钱树、.*?)等质量等级划分标准", content, re.S)
  if not match:
    raise ValueError("Horticulture proposal section changed.")
  names = list(dict.fromkeys(compact(match[1]).split("、")))
  return [{
    "recordId": f"F06:proposed:{i}", "sourceId": "F06",
    "sourceWork": "standard-project:20242424-T-432", "sourceRow": i,
    "sourceSection": "范围和主要技术内容/拟新增产品",
    "sourceProvince": "非地区性项目资料", "sourceRegion": "未核验",
    "chineseNameRaw": name, "scientificNameRaw": "", "sourceRank": None,
    "sourceType": "horticulture-standard-proposal", "frequency": "未明确",
    "noteRaw": "", "suitableUses": ["盆栽产品候选"],
    "evidence": "B-proposed-product-scope-not-final-standard-or-market-count",
  } for i, name in enumerate(names, 1)]


def extract_indoor(cache, source):
  content = source_html_body(cache, source)
  match = re.search(
    r"([\u4e00-\u9fff]+(?:、[\u4e00-\u9fff]+)+)等都是常见的室内栽?观赏植物", content,
  )
  if not match or len(match[1].split("、")) != 4:
    raise ValueError("Indoor-use source paragraph changed.")
  return [{
    "recordId": f"F26:indoor:{i}", "sourceId": "F26",
    "sourceWork": source["sourceWork"], "sourceRow": i,
    "sourceSection": "天南星科介绍/常见室内栽观赏植物",
    "sourceProvince": "机构室内用途资料", "sourceRegion": "未核验",
    "chineseNameRaw": name, "scientificNameRaw": "", "sourceRank": None,
    "sourceType": "institutional-indoor-common-name-list", "frequency": "未明确",
    "noteRaw": "", "suitableUses": ["室内观赏"],
    "evidence": "B-explicit-common-indoor-use-not-regional-abundance",
  } for i, name in enumerate(match[1].split("、"), 1)]


def load_catalog(path):
  by_name = defaultdict(list)
  by_chinese = defaultdict(list)
  with sqlite3.connect(f"file:{path.resolve()}?mode=ro", uri=True) as connection:
    for raw, in connection.execute("SELECT data FROM species"):
      row = json.loads(raw)
      item = {
        "speciesId": row["id"], "scientificName": row["scientificName"],
        "chineseName": row.get("chineseName"), "isVisible": row.get("isVisible", False),
      }
      by_name[row["scientificName"]].append(item)
      for name in {row.get("chineseName"), row.get("commonName"),
                   *row.get("chineseAliases", [])} - {None, ""}:
        by_chinese[name].append(item)
  return dict(by_name), dict(by_chinese)


def classify_name(row, by_compact):
  raw = row["scientificNameRaw"]
  declared = row.get("sourceRank")
  if declared in {"品种", "变种", "亚种", "变型"}:
    return {"rank": declared, "key": None, "name": text(raw), "match": None,
            "method": "source-rank-retained"}
  if re.search(r"\b(?:var|subsp|ssp|f)\.\s*[a-z]", raw):
    return {"rank": "种下等级", "key": None, "name": text(raw), "match": None,
            "method": "explicit-infraspecific-name"}
  if any(marker in raw for marker in ("'", "’", "‘", '"')):
    return {"rank": "品种或存疑", "key": None, "name": text(raw), "match": None,
            "method": "quoted-name-not-collapsed"}
  if "×" in raw or re.search(r"\b[A-Z][a-z]+\s+x\s+", raw):
    return {"rank": "杂交待核验", "key": "hybrid:" + compact(raw),
            "name": text(raw), "match": None, "method": "hybrid-scope-review"}
  joined = compact(raw)
  matches = []
  for end in range(3, min(len(joined), 100) + 1):
    rest = joined[end:]
    if joined[:end] in by_compact and (not rest or rest[0].isupper() or rest[0] == "("):
      matches.extend(by_compact[joined[:end]])
  if len(matches) == 1:
    item = matches[0]
    return {"rank": "种", "key": "catalog:" + item["speciesId"],
            "name": item["scientificName"], "match": item,
            "method": "exact-scientific-label-with-whitespace-author-normalization"}
  latin = re.match(r"^([A-Z][a-z]+)\s+([a-z][a-z-]+)(?:\s|$)", text(raw))
  if latin:
    name = " ".join(latin.groups())
    return {"rank": "种级候选待核验", "key": "unmatched-literal:" + name, "name": name,
            "match": None, "method": "literal-binomial-not-synonym-resolved"}
  return {"rank": "待核验", "key": "unresolved:" + row["recordId"],
          "name": text(raw), "match": None, "method": "unresolved-source-label"}


def make_pool(records, by_name, by_chinese):
  by_compact = defaultdict(list)
  for name, values in by_name.items():
    if re.fullmatch(r"[A-Z][a-z]+ [a-z][a-z-]+", name):
      by_compact[compact(name)].extend(values)
  pool, excluded = {}, []
  for row in records:
    parsed = classify_name(row, by_compact)
    row["nameResolution"] = parsed
    if parsed["key"] is None:
      excluded.append(row["recordId"])
      continue
    key = parsed["key"]
    if key not in pool:
      pool[key] = {
        "candidateId": "review:" + hashlib.sha256(key.encode()).hexdigest()[:20],
        "scientificLabel": parsed["name"], "rankStatus": parsed["rank"],
        "catalogMatch": parsed["match"], "sourceNames": [],
        "sourceRecordIds": [], "sourceWorks": set(), "sourceRegions": set(),
        "frequencies": Counter(), "sourceTypes": set(), "urbanRecommendationWorks": set(),
        "nameOnlyCatalogHints": [], "trainingApproved": False,
        "selectionStatus": "pending-evidence-taxonomy-and-regional-review",
      }
    candidate = pool[key]
    candidate["sourceNames"].append(row["chineseNameRaw"])
    candidate["sourceRecordIds"].append(row["recordId"])
    candidate["sourceWorks"].add(row["sourceWork"])
    candidate["sourceRegions"].add(row["sourceRegion"])
    candidate["frequencies"][row["frequency"]] += 1
    candidate["sourceTypes"].add(row["sourceType"])
    if "城市" in row["suitableUses"]:
      candidate["urbanRecommendationWorks"].add(row["sourceWork"])
    if not raw_name_present(row):
      # Hints only: Chinese product names are never silently merged into species.
      name = re.split("[（(]", row["chineseNameRaw"])[0]
      candidate["nameOnlyCatalogHints"] = by_chinese.get(name, [])
  for candidate in pool.values():
    candidate["sourceNames"] = list(dict.fromkeys(candidate["sourceNames"]))
    for field in ("sourceWorks", "sourceRegions", "sourceTypes", "urbanRecommendationWorks"):
      candidate[field] = sorted(candidate[field])
    candidate["frequencies"] = dict(candidate["frequencies"])
    candidate["priorityReview"] = bool(
      set(candidate["frequencies"]) & STRONG_FREQUENCIES
      or candidate["urbanRecommendationWorks"]
      or "horticulture-standard-proposal" in candidate["sourceTypes"]
      or "institutional-indoor-common-name-list" in candidate["sourceTypes"]
    )
  return sorted(pool.values(), key=lambda row: row["candidateId"]), excluded


def raw_name_present(row):
  return bool(row["scientificNameRaw"].strip())


def export_csv(path, candidates):
  with path.open("w", encoding="utf-8-sig", newline="") as stream:
    writer = csv.writer(stream)
    writer.writerow(["候选ID（非排名）", "来源中文名", "学名标签", "分类状态", "本地目录ID候选",
                     "来源地区（非全国分布）", "独立资料数", "原始常见度标记",
                     "城市绿化推荐资料数", "优先核对", "训练已批准", "来源记录定位"])
    for row in candidates:
      writer.writerow([
        row["candidateId"], " / ".join(row["sourceNames"]), row["scientificLabel"],
        row["rankStatus"], (row["catalogMatch"] or {}).get("speciesId", ""),
        " / ".join(row["sourceRegions"]), len(row["sourceWorks"]),
        json.dumps(row["frequencies"], ensure_ascii=False),
        len(row["urbanRecommendationWorks"]), "是" if row["priorityReview"] else "否", "否",
        " | ".join(row["sourceRecordIds"]),
      ])


def compare_seeds(seeds, pool):
  rows = []
  for seed in seeds:
    ids = {item["speciesId"] for item in seed["catalogNameMatches"]}
    by_taxon = [
      item for item in pool if item["catalogMatch"] and item["catalogMatch"]["speciesId"] in ids
    ]
    by_name = [
      item for item in pool if any(
        re.split("[（(]", name)[0] == seed["proposedName"] for name in item["sourceNames"]
      )
    ]
    all_matches = {item["candidateId"]: item for item in by_taxon + by_name}
    rows.append({
      "seedId": seed["candidateId"], "name": seed["proposedName"],
      "scene": seed["proposedScene"],
      "status": "taxon-label-found" if by_taxon else "source-name-only" if by_name else "gap",
      "candidateIds": sorted(all_matches),
      "hasRegionalFrequencyEvidence": any(
        set(item["frequencies"]) & STRONG_FREQUENCIES for item in all_matches.values()
      ),
      "hasExplicitIndoorUseEvidence": any(
        "institutional-indoor-common-name-list" in item["sourceTypes"]
        for item in all_matches.values()
      ),
      "trainingApproved": False,
    })
  return rows


def build(cache, registry_path, catalog_path, output, seeds_path=SEEDS):
  if output.exists():
    raise ValueError("Output exists; choose a new directory. Inputs are never overwritten.")
  sources = {row["id"]: row for row in json.loads(registry_path.read_text())["sources"]}
  records = []
  for identifier, config in PDF_CONFIG.items():
    records.extend(extract_pdf(cache, sources[identifier], config))
  records.extend(extract_workbook(cache, sources["F24"]))
  records.extend(extract_horticulture(cache, sources["F06"]))
  records.extend(extract_indoor(cache, sources["F26"]))
  if len({row["recordId"] for row in records}) != len(records):
    raise ValueError("Duplicate source record identities.")
  names, chinese = load_catalog(catalog_path)
  pool, excluded = make_pool(records, names, chinese)
  priority = [row for row in pool if row["priorityReview"]]
  seed_coverage = compare_seeds(json.loads(seeds_path.read_text())["candidates"], pool)
  report = {
    "createdAt": datetime.now(UTC).isoformat(),
    "policy": "source-review-pool-v1-not-training-selection",
    "rawRecordCount": len(records), "candidateLabelCount": len(pool),
    "priorityReviewLabelCount": len(priority), "approvedTrainingClasses": 0,
    "excludedInfraspecificCultivarRecords": len(excluded),
    "excludedRecordIds": excluded,
    "bySource": dict(Counter(row["sourceId"] for row in records)),
    "bySourceType": dict(Counter(row["sourceType"] for row in records)),
    "rankStatuses": dict(Counter(row["rankStatus"] for row in pool)),
    "catalogMatchedCandidates": sum(row["catalogMatch"] is not None for row in pool),
    "hiddenCatalogMatchesRetained": sum(
      row["catalogMatch"] is not None and not row["catalogMatch"]["isVisible"] for row in pool
    ),
    "sourceRegionRecordCounts": dict(Counter(row["sourceRegion"] for row in records)),
    "sourceTypeFrequencyCounts": {
      kind: dict(Counter(row["frequency"] for row in records if row["sourceType"] == kind))
      for kind in sorted({row["sourceType"] for row in records})
    },
    "priorityRanks": dict(Counter(row["rankStatus"] for row in priority)),
    "priorityOnlyShanghai": sum(row["sourceRegions"] == ["华东"] for row in priority),
    "seedCoverage": dict(Counter(row["status"] for row in seed_coverage)),
    "seedGaps": [row["name"] for row in seed_coverage if row["status"] == "gap"],
    "seedSnapshotSha256": digest(seeds_path),
    "catalogSha256": digest(catalog_path),
    "sources": [{
      "id": identifier, "title": sources[identifier]["title"],
      "url": sources[identifier]["url"],
      "sha256": sources[identifier]["directFetchCheck"]["sha256"],
    } for identifier in ("F18", "F19", "F20", "F24", "F06", "F26")],
    "limitations": [
      "Candidate labels, not verified accepted species or a national commonness ranking.",
      "Shanghai survey dominates this batch. 华南/西南/西北 source coverage still missing.",
      "Native suitability is not actual planting abundance.",
      "Product proposal labels are not a final standard or resolved species.",
      "No fuzzy synonyms, manual alias database writes, private records, or training.",
      "Cultivars/infraspecific ranks kept in raw records, not collapsed into parent species.",
      "Readable source metadata does not establish redistribution or photo-training rights.",
      "Priority review includes 多见 separately; it does not mean all are 常见.",
    ],
  }
  output.mkdir(parents=True)
  with (output / "source-records.jsonl").open("w", encoding="utf-8") as stream:
    for row in records:
      stream.write(json.dumps(row, ensure_ascii=False) + "\n")
  write_json(output / "candidates.json", pool)
  write_json(output / "coverage.json", report)
  write_json(output / "seed-coverage.json", seed_coverage)
  with (output / "seed-coverage.csv").open("w", encoding="utf-8-sig", newline="") as stream:
    writer = csv.writer(stream)
    writer.writerow(["草案称呼", "场景", "对齐状态", "存在地区常见或多见证据",
                     "有明确室内用途证据", "候选ID", "批准训练"])
    for row in seed_coverage:
      writer.writerow([row["name"], row["scene"], row["status"],
                       row["hasRegionalFrequencyEvidence"], row["hasExplicitIndoorUseEvidence"],
                       " | ".join(row["candidateIds"]), "否"])
  export_csv(output / "candidate-pool.csv", pool)
  export_csv(output / "priority-review.csv", priority)
  print(json.dumps({key: report[key] for key in (
    "rawRecordCount", "candidateLabelCount", "priorityReviewLabelCount",
    "excludedInfraspecificCultivarRecords", "catalogMatchedCandidates",
    "bySource", "rankStatuses", "sourceTypeFrequencyCounts",
  )}, ensure_ascii=False, indent=2))
  return report


def main():
  parser = argparse.ArgumentParser(description=__doc__)
  parser.add_argument("--cache", type=Path, default=CACHE)
  parser.add_argument("--registry", type=Path, default=REGISTRY)
  parser.add_argument("--catalog", type=Path, default=CATALOG)
  parser.add_argument("--seeds", type=Path, default=SEEDS)
  parser.add_argument("--output", type=Path, required=True)
  args = parser.parse_args()
  build(args.cache, args.registry, args.catalog, args.output, args.seeds)


if __name__ == "__main__":
  main()
